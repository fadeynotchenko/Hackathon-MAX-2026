"""Сборка DOCX из готового текста документа.

Тело встроенных шаблонов — текст с подставленными значениями, поэтому DOCX
собирается программно: абзац на строку, моноширинного форматирования нет.
Когда появятся шаблоны-файлы компании, рядом встанет второй сборщик — он будет
подставлять значения в готовый .docx, сохраняя её оформление.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from docx.shared import Pt

BASE_FONT = "Times New Roman"
BASE_SIZE_PT = 12


def build_docx(title: str, text: str) -> bytes:
    """Заголовок — название документа, если текст не начинается с него сам.

    У встроенных шаблонов первая строка текста и есть заголовок («Счёт на оплату
    № 17 от …»): название «Счёт на оплату» над ней повторяло бы её. Тогда
    заголовком становится сама первая строка."""
    document = DocxDocument()
    document.core_properties.title = title
    style = document.styles["Normal"]
    style.font.name = BASE_FONT
    style.font.size = Pt(BASE_SIZE_PT)

    lines = text.splitlines()
    first = next((index for index, line in enumerate(lines) if line.strip()), None)
    title_in_text = first is not None and bool(title) and lines[first].startswith(title)
    if title and not title_in_text:
        _add_heading(document, title)

    for index, line in enumerate(lines):
        if title_in_text and index == first:
            _add_heading(document, line)
        else:
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_heading(document: DocxDocument, text: str) -> None:
    run = document.add_paragraph().add_run(text)
    run.bold = True
    run.font.size = Pt(BASE_SIZE_PT + 2)
