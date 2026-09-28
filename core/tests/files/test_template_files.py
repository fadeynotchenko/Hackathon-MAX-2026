"""Файл-образец шаблона: текст из DOCX и PDF, сборка документа в оформлении DOCX."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import parse_xml

from core.domain.places import Place, apply_places, label_of, place_spans
from core.files import (
    TemplateFileError,
    docx_layout,
    docx_lines,
    fill_docx,
    mark_blank_cells,
    pdf_lines,
)
from tests.samples import offer_docx, text_pdf


def test_place_spans_prefer_longer_places_and_keep_context() -> None:
    line = "Заказчик: ООО «Альфа Плюс», исполнитель: ________, подпись: ________"
    places = [
        ("client_short", Place("Альфа")),
        ("client_name", Place("ООО «Альфа Плюс»")),
        ("seller_director", Place("________", before="исполнитель: ")),
    ]
    assert apply_places(line, places) == (
        "Заказчик: {{client_name}}, исполнитель: {{seller_director}}, подпись: ________"
    )
    assert place_spans("нет совпадений", places) == []
    assert label_of("{{  Название   клиента }}") == "Название клиента"


def test_docx_lines_read_header_body_and_tables() -> None:
    assert docx_lines(offer_docx()) == [
        "ООО «Мастер» · ИНН 7707083893",
        "Коммерческое предложение",
        "Для: ООО «Альфа»",
        "Работы: монтаж оборудования",
        "Стоимость",
        "180 000 руб.",
        "Заказчик: ________\tИсполнитель: ________",
        "Типовые условия: оплата в течение 14 дней.",
    ]


def test_fill_docx_keeps_formatting_of_the_sample() -> None:
    places = [
        ("seller_name", Place("ООО «Мастер»")),
        ("client_name", Place("ООО «Альфа»")),
        ("total", Place("180 000")),
        ("client_director", Place("________", before="Заказчик: ")),
        ("works", Place("монтаж оборудования")),
    ]
    filled = fill_docx(
        offer_docx(),
        places=places,
        context={
            "seller_name": "ИП Иванов",
            "client_name": "ООО «Бета»",
            "total": "250 000,00",
            "works": "Монтаж\nПусконаладка",
        },
        blank="__________",
    )

    assert docx_lines(filled) == [
        "ИП Иванов · ИНН 7707083893",
        "Коммерческое предложение",
        "Для: ООО «Бета»",
        "Работы: Монтаж\nПусконаладка",
        "Стоимость",
        "250 000,00 руб.",
        "Заказчик: __________\tИсполнитель: ________",
        "Типовые условия: оплата в течение 14 дней.",
    ]
    client = Document(BytesIO(filled)).paragraphs[1]
    assert [(run.text, run.bold) for run in client.runs if run.text] == [
        ("Для: ", None),
        ("ООО «Бета»", True),
    ], "значение берёт оформление фрагмента, который заменило"


def test_fill_docx_replaces_marks_written_in_the_file() -> None:
    places = [("client_name", Place("{{ Название клиента }}")), ("total", Place("{{Сумма}}"))]
    filled = fill_docx(
        offer_docx(marked=True), places=places, context={"total": "5,00"}, blank="____"
    )
    lines = docx_lines(filled)
    assert "Для: ____" in lines and "5,00 руб." in lines


def test_unreadable_samples_are_rejected() -> None:
    with pytest.raises(TemplateFileError):
        docx_lines(b"PK\x03\x04 not really a docx")
    with pytest.raises(TemplateFileError):
        pdf_lines(text_pdf([]))
    assert pdf_lines(text_pdf(["Invoice for ACME Corp", "Total: 100 USD"]))[:2] == [
        "Invoice for ACME Corp",
        "Total: 100 USD",
    ]


def _blank_form() -> bytes:
    """Бланк, где места — пустые ячейки с линейкой: «ИНН [____]», подстрочник под ней."""
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ИНН"
    table.cell(1, 1).text = "цифрами"
    border = (
        '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:bottom w:val="single" w:sz="4" w:color="000000"/></w:tcBorders>'
    )
    table.cell(0, 1)._tc.get_or_add_tcPr().append(parse_xml(border))
    document.add_paragraph("Итого: {{total|words}}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_blank_cells_get_label_marks_and_stay_empty_without_value() -> None:
    marked = mark_blank_cells(_blank_form())
    assert marked is not None
    assert "{{ИНН (цифрами)}}" in docx_lines(marked)
    assert docx_layout(marked)[0] == "ИНН  {{ИНН (цифрами)}}", "строка таблицы — одна строка"
    assert mark_blank_cells(offer_docx()) is None, "линеек-ячеек нет — файл не трогаем"

    places = [("inn", Place("{{ИНН (цифрами)}}"))]
    empty = fill_docx(marked, places=places, context={}, blank="__________")
    assert docx_lines(empty)[:2] == ["ИНН", ""], "пустое значение — пустая ячейка, как в бланке"
    assert "Итого: __________" in docx_lines(empty), "в тексте пустое — прочерк"

    filled = fill_docx(
        marked,
        places=places,
        context={"inn": "7707083893", "total|words": "Сто рублей 00 копеек"},
        blank="__________",
    )
    assert docx_lines(filled)[1] == "7707083893"
    assert "Итого: Сто рублей 00 копеек" in docx_lines(filled), "вариант записи из маркера"


def _conditional_form() -> bytes:
    """Бланк с условными кусками: в строке, целым абзацем, в строке таблицы
    «подпись | значение» и в таблице позиций с общей ячейкой срока."""
    document = Document()
    supplier = document.add_paragraph("Поставщик: {{name}}[[, ")
    supplier.add_run("КПП").bold = True
    supplier.add_run(" {{kpp}}]][[, {{address}}]]")
    document.add_paragraph("[[Порядок оплаты: {{terms}}.]]")
    document.add_paragraph("[[Тел. {{phone}}]][[, почта {{email}}]]")
    requisites = document.add_table(rows=2, cols=2)
    requisites.cell(0, 0).text = "Наименование"
    requisites.cell(0, 1).text = "{{name}}"
    requisites.cell(1, 0).text = "Адрес"
    requisites.cell(1, 1).text = "[[{{address}}]]"
    items = document.add_table(rows=2, cols=4)
    for cell, text in zip(items.rows[0].cells, ("№", "Наименование", "Сумма", "Срок"), strict=True):
        cell.text = text
    for cell, text in zip(
        items.rows[1].cells,
        ("{{items.n}}", "{{items.name}}", "{{items.amount}}", "{{term}}"),
        strict=True,
    ):
        cell.text = text
    document.add_paragraph("Всего наименований {{items|count}}, на сумму {{items|sum}}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_conditional_pieces_vanish_with_labels_and_rows() -> None:
    empty = fill_docx(
        _conditional_form(), places=[], context={"name": "ООО «Ромашка»"}, blank="____"
    )
    opened = Document(BytesIO(empty))
    assert [p.text for p in opened.paragraphs][:2] == [
        "Поставщик: ООО «Ромашка»",
        "Всего наименований ____, на сумму ____",
    ], "абзацы из одних подписей убраны, пустое вне кусков — линия"
    requisites, items = opened.tables
    assert [row.cells[0].text for row in requisites.rows] == ["Наименование"], (
        "строка «Адрес» без адреса ушла целиком"
    )
    assert [cell.text for cell in items.rows[1].cells] == ["", "", "", ""], (
        "пустой список — одна пустая строка для записи от руки"
    )
    assert "[[" not in "\n".join(docx_lines(empty))

    full = fill_docx(
        _conditional_form(),
        places=[],
        context={
            "name": "ООО «Ромашка»",
            "kpp": "772801001",
            "email": "a@b.example",
            "address": "Москва",
        },
        blank="____",
    )
    paragraphs = Document(BytesIO(full)).paragraphs
    assert paragraphs[0].text == "Поставщик: ООО «Ромашка», КПП 772801001, Москва"
    assert [(run.text, run.bold) for run in paragraphs[0].runs if run.text][1] == ("КПП", True), (
        "скобки убраны, оформление кусков осталось"
    )
    assert paragraphs[1].text == "почта a@b.example", "запятая в начале не остаётся"


def test_items_row_is_repeated_and_shared_cell_is_merged() -> None:
    context = {
        "items|count": "3",
        "items|sum": "600,00",
        "term": "10 дней",
        **{f"items.{i}.n": str(i) for i in (1, 2, 3)},
        **{f"items.{i}.name": f"Позиция {i}" for i in (1, 2, 3)},
        **{f"items.{i}.amount": f"{i}00,00" for i in (1, 2, 3)},
    }
    filled = fill_docx(_conditional_form(), places=[], context=context, blank="____")
    opened = Document(BytesIO(filled))
    table = opened.tables[1]
    assert [[cell.text for cell in row.cells[:3]] for row in table.rows[1:]] == [
        ["1", "Позиция 1", "100,00"],
        ["2", "Позиция 2", "200,00"],
        ["3", "Позиция 3", "300,00"],
    ]
    merges = [
        row._tr.tc_lst[3].tcPr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}vMerge"
        )
        for row in table.rows[1:]
    ]
    assert [
        m.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") for m in merges
    ] == ["restart", None, None], "срок — одна ячейка на все позиции"
    assert table.rows[1].cells[3].text == "10 дней"
    properties = [child.tag.rsplit("}", 1)[1] for child in table.rows[1]._tr.tc_lst[3].tcPr]
    assert properties.index("vMerge") == properties.index("tcW") + 1, "порядок свойств — как у Word"
    assert "Всего наименований 3, на сумму 600,00" in [p.text for p in opened.paragraphs]
