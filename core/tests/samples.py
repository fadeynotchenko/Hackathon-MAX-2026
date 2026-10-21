"""Файлы-образцы для тестов: DOCX с оформлением, PDF с текстом и голосовые."""

from __future__ import annotations

import math
import struct
from io import BytesIO

import av
from docx import Document


def offer_docx(*, marked: bool = False) -> bytes:
    """КП «Мастера» для «Альфы»: шапка с реквизитами, название клиента жирным
    и порезано Word'ом на два куска, сумма — в таблице."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "ООО «Мастер» · ИНН 7707083893"
    document.add_paragraph("Коммерческое предложение")
    paragraph = document.add_paragraph("Для: ")
    if marked:
        paragraph.add_run("{{ Название клиента }}").bold = True
    else:
        paragraph.add_run("ООО «Аль").bold = True
        paragraph.add_run("фа»").bold = True
    document.add_paragraph("Работы: монтаж оборудования")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Стоимость"
    table.cell(0, 1).text = "{{Сумма}} руб." if marked else "180 000 руб."
    document.add_paragraph("Заказчик: ________\tИсполнитель: ________")
    document.add_paragraph("Типовые условия: оплата в течение 14 дней.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# Ответ помощника на offer_docx(): клиент и сумма — места для данных.
OFFER_PLACES = {
    "title": "Коммерческое предложение",
    "kind": "offer",
    "places": [
        {
            "text": "ООО «Альфа»",
            "before": "Для: ",
            "label": "Название клиента",
            "type": "text",
            "key": "client_name",
        },
        {"text": "180 000", "before": "", "label": "Стоимость", "type": "money", "key": "total"},
    ],
}


def text_pdf(lines: list[str]) -> bytes:
    """Одностраничный PDF с текстовым слоем (Helvetica, латиница)."""
    stream = "BT /F1 12 Tf 72 720 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines)
    stream += " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        " /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return out


def voice_clip(container: str = "ogg", codec: str = "libopus", *, seconds: float = 0.5) -> bytes:
    """Настоящее голосовое заданного формата: тон 440 Гц. Байты с одной сигнатурой
    не годятся — голосовое перед моделью декодируется."""
    rate = 48_000
    buffer = BytesIO()
    with av.open(buffer, "w", format=container) as output:
        stream = output.add_stream(codec, rate=rate, layout="mono")
        chunk = stream.codec_context.frame_size or 1024
        total = int(rate * seconds)
        for start in range(0, total, chunk):
            count = min(chunk, total - start)
            samples = [
                int(8000 * math.sin(2 * math.pi * 440 * (start + i) / rate)) for i in range(count)
            ]
            frame = av.AudioFrame(format="s16", layout="mono", samples=count)
            frame.planes[0].update(struct.pack(f"<{count}h", *samples))
            frame.rate = rate
            frame.pts = start
            for packet in stream.encode(frame):
                output.mux(packet)
        for packet in stream.encode(None):
            output.mux(packet)
    return buffer.getvalue()
